import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

import docx
import openpyxl
import pymupdf

from app.services.document_parser import ParseError, parse_document


def rejects(data: bytes, ext: str) -> bool:
    try:
        parse_document(data, ext)
        return False
    except ParseError:
        return True


# TXT and Markdown
txt = parse_document(b"Synthetic meeting notes", ".txt")
print("txt works:", "meeting notes" in txt.pages[0].text)
md = parse_document(b"# Title\n\nSynthetic budget plan", ".md")
print("md works:", "budget plan" in md.pages[0].text)

# PDF with two pages: page numbers must be kept
pdf = pymupdf.open()
pdf.new_page().insert_text((72, 72), "First page synthetic text")
pdf.new_page().insert_text((72, 72), "Second page synthetic text")
parsed_pdf = parse_document(pdf.tobytes(), ".pdf")
print("pdf page numbers kept:", [p.page_number for p in parsed_pdf.pages] == [1, 2])
print("pdf text extracted:", "Second page" in parsed_pdf.pages[1].text)

# DOCX with a paragraph and a table
document = docx.Document()
document.add_paragraph("Synthetic project report")
table = document.add_table(rows=1, cols=2)
table.cell(0, 0).text = "Expense"
table.cell(0, 1).text = "Travel"
buffer = io.BytesIO()
document.save(buffer)
parsed_docx = parse_document(buffer.getvalue(), ".docx")
print("docx paragraph and table:",
      "project report" in parsed_docx.pages[0].text and "Travel" in parsed_docx.pages[0].text)

# XLSX
workbook = openpyxl.Workbook()
workbook.active.title = "Budget"
workbook.active.append(["Item", "Cost"])
workbook.active.append(["Rent", 1200])
buffer = io.BytesIO()
workbook.save(buffer)
parsed_xlsx = parse_document(buffer.getvalue(), ".xlsx")
print("xlsx works:", "Rent" in parsed_xlsx.pages[0].text and "Budget" in parsed_xlsx.pages[0].text)

# Bad input must fail cleanly
print("corrupt pdf rejected:", rejects(b"this is not a pdf", ".pdf"))
print("corrupt docx rejected:", rejects(b"this is not a docx", ".docx"))
print("unsupported type rejected:", rejects(b"data", ".exe"))
print("empty text rejected:", rejects(b"   ", ".txt"))
print("binary disguised as text rejected:", rejects(b"abc\x00def", ".txt"))