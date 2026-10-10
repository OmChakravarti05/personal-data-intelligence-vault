import io
from dataclasses import dataclass

import docx
import openpyxl
import pymupdf

MAX_PAGES = 500
MAX_CHARS = 5_000_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx"}


class ParseError(Exception):
    """The file could not be read. The message is safe to show to the user."""


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int | None  # real page numbers for PDFs, None for other formats
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[ExtractedPage]
    page_count: int | None


def parse_document(data: bytes, extension: str) -> ExtractedDocument:
    extension = extension.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"Unsupported file type: {extension or 'none'}")
    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".xlsx": _parse_xlsx,
    }
    try:
        result = parsers[extension](data)
    except ParseError:
        raise
    except Exception as exc:  # corrupted files raise many different library errors
        raise ParseError("The file is corrupted or could not be read") from exc

    total = sum(len(page.text) for page in result.pages)
    if total > MAX_CHARS:
        raise ParseError("The document is too large to process")
    if not any(page.text.strip() for page in result.pages):
        raise ParseError("No extractable text found (the file may be empty or scanned)")
    return result


def _parse_pdf(data: bytes) -> ExtractedDocument:
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        if doc.needs_pass:
            raise ParseError("Password-protected PDFs are not supported")
        if doc.page_count > MAX_PAGES:
            raise ParseError(f"PDF has more than {MAX_PAGES} pages")
        pages = [
            ExtractedPage(index + 1, page.get_text())
            for index, page in enumerate(doc)
        ]
        return ExtractedDocument(pages, len(pages))
    finally:
        doc.close()


def _parse_docx(data: bytes) -> ExtractedDocument:
    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append("\t".join(cells))
    return ExtractedDocument([ExtractedPage(None, "\n".join(lines))], None)


def _parse_text(data: bytes) -> ExtractedDocument:
    if b"\x00" in data:
        raise ParseError("The file looks like binary data, not text")
    text = data.decode("utf-8-sig", errors="replace")
    return ExtractedDocument([ExtractedPage(None, text)], None)


def _parse_xlsx(data: bytes) -> ExtractedDocument:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        pages = []
        for sheet in workbook.worksheets:
            lines = [f"Sheet: {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    lines.append("\t".join(cells))
            pages.append(ExtractedPage(None, "\n".join(lines)))
        return ExtractedDocument(pages, None)
    finally:
        workbook.close()